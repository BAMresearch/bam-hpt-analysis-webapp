"""T4.2 summary export (C2 / C2b) — schema t42_summary_v1.

Builds one row per results entry for CSV / Excel hand-off.
Column contract: HPT/WP4/T4.2_Evidence/EXPORT_SCHEMA_T42.md

C2b: optional column subset at download time; local default prefs file.
"""
from __future__ import annotations

import csv
import json
import os
import sqlite3
import subprocess
from datetime import datetime, timezone
from io import BytesIO, StringIO
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

SCHEMA_ID = "t42_summary_v1"

# Frozen column order (C1 / C2). Do not reorder without a schema version bump.
SUMMARY_COLUMNS: Tuple[str, ...] = (
    "tool_version",
    "entry_id",
    "file_name",
    "data_set",
    "lab_id",
    "hp_id",
    "test_label",
    "test_cond",
    "flow_config",
    "profile_id",
    "condition_set_id",
    "appliance_type",
    "cop_dataset",
    "start_time_s",
    "end_time_s",
    "cycle_type",
    "avg_t_db",
    "avg_t_wb",
    "avg_t_supply",
    "avg_t_return_emu",
    "avg_t_return_calc",
    "DTret",
    "avg_volume_flow",
    "avg_mass_flow",
    "avg_Q_corr",
    "avg_P_corr",
    "COP_corr",
    "avg_Q_uncorr",
    "avg_P_uncorr",
    "COP_uncorr",
    "avg_Q_bam_corr",
    "avg_P_bam_corr",
    "COP_bam_corr",
    "dev_test_condition",
    "dev_db_setpoint",
    "dev_wb_setpoint",
    "dev_tsup_setpoint",
    "dev_total_points",
    "dev_db_outside_count",
    "dev_db_percentage",
    "dev_wb_outside_count",
    "dev_wb_percentage",
    "dev_tsup_outside_count",
    "dev_tsup_percentage",
    "dev_dtreturn_outside_count",
    "dev_dtreturn_percentage",
    "dev_flow_outside_count",
    "dev_flow_percentage",
    "dev_mean_flow_dev_pct",
    "dev_max_db_pos",
    "dev_max_db_neg",
    "dev_max_wb_pos",
    "dev_max_wb_neg",
    "dev_max_tsup_pos",
    "dev_max_tsup_neg",
    "dev_max_dtreturn_pos",
    "dev_max_dtreturn_neg",
    "dev_max_flow_pos_pct",
    "dev_max_flow_neg_pct",
    "pass_instantaneous",
    "pass_mean",
    "pass_controllability",
    "notes",
    # Optional monitoring means (air-to-water REV1), appended in v0.1.0.
    # Append-only: existing columns keep their position, so t42_summary_v1 readers that
    # match on header name (not index) are unaffected. Empty when the lab did not supply
    # the series — never 0.
    "avg_lab_air_db",
    "avg_lab_atm_pressure",
    "avg_cp_out_sink",
    "avg_cp_in_sink",
    "avg_density_sink",
    "avg_compressor_frequency",
    "avg_fan_speed",
    "avg_voltage",
    "avg_current",
    "avg_frequency",
    "avg_cos_phi",
    "avg_real_buh_power",
    "avg_t_in_cond",
    "avg_t_out_cond",
    "avg_t_in_evap",
    "avg_t_out_evap",
)

# Export names that are also the DB column names (optional monitoring means).
OPTIONAL_MONITORING_COLUMNS: Tuple[str, ...] = (
    "avg_lab_air_db",
    "avg_lab_atm_pressure",
    "avg_cp_out_sink",
    "avg_cp_in_sink",
    "avg_density_sink",
    "avg_compressor_frequency",
    "avg_fan_speed",
    "avg_voltage",
    "avg_current",
    "avg_frequency",
    "avg_cos_phi",
    "avg_real_buh_power",
    "avg_t_in_cond",
    "avg_t_out_cond",
    "avg_t_in_evap",
    "avg_t_out_evap",
)

_APP_DIR = os.path.dirname(os.path.abspath(__file__))
_PREFS_PATH = os.path.join(_APP_DIR, "config", "export_summary_prefs.json")

# UI grouping for advanced column picker (C2b). Order of groups is display-only;
# export order always follows SUMMARY_COLUMNS.
COLUMN_GROUPS: Tuple[Tuple[str, Tuple[str, ...]], ...] = (
    (
        "Identity & binding",
        (
            "tool_version",
            "entry_id",
            "file_name",
            "data_set",
            "lab_id",
            "hp_id",
            "test_label",
            "test_cond",
            "flow_config",
            "profile_id",
            "condition_set_id",
            "appliance_type",
            "cop_dataset",
        ),
    ),
    (
        "Window & means / COP",
        (
            "start_time_s",
            "end_time_s",
            "cycle_type",
            "avg_t_db",
            "avg_t_wb",
            "avg_t_supply",
            "avg_t_return_emu",
            "avg_t_return_calc",
            "DTret",
            "avg_volume_flow",
            "avg_mass_flow",
            "avg_Q_corr",
            "avg_P_corr",
            "COP_corr",
            "avg_Q_uncorr",
            "avg_P_uncorr",
            "COP_uncorr",
            "avg_Q_bam_corr",
            "avg_P_bam_corr",
            "COP_bam_corr",
        ),
    ),
    (
        "Deviations",
        (
            "dev_test_condition",
            "dev_db_setpoint",
            "dev_wb_setpoint",
            "dev_tsup_setpoint",
            "dev_total_points",
            "dev_db_outside_count",
            "dev_db_percentage",
            "dev_wb_outside_count",
            "dev_wb_percentage",
            "dev_tsup_outside_count",
            "dev_tsup_percentage",
            "dev_dtreturn_outside_count",
            "dev_dtreturn_percentage",
            "dev_flow_outside_count",
            "dev_flow_percentage",
            "dev_mean_flow_dev_pct",
            "dev_max_db_pos",
            "dev_max_db_neg",
            "dev_max_wb_pos",
            "dev_max_wb_neg",
            "dev_max_tsup_pos",
            "dev_max_tsup_neg",
            "dev_max_dtreturn_pos",
            "dev_max_dtreturn_neg",
            "dev_max_flow_pos_pct",
            "dev_max_flow_neg_pct",
        ),
    ),
    (
        "Optional monitoring (A/W)",
        OPTIONAL_MONITORING_COLUMNS,
    ),
    (
        "Notes",
        ("notes",),
    ),
)

# Kept in SUMMARY_COLUMNS for a stable future hand-off layout, but hidden from the
# Export UI and omitted from the recommended default until pass/fail rules exist.
DEFERRED_PASS_COLUMNS: Tuple[str, ...] = (
    "pass_instantaneous",
    "pass_mean",
    "pass_controllability",
)

# Default tick-set / “all usual columns” (excludes deferred empty pass_* placeholders).
DEFAULT_EXPORT_COLUMNS: Tuple[str, ...] = tuple(
    c for c in SUMMARY_COLUMNS if c not in DEFERRED_PASS_COLUMNS
)

# DB column candidates → export name (first match wins; case-insensitive scan)
_DB_FIELD_CANDIDATES: Dict[str, Tuple[str, ...]] = {
    "file_name": ("file_name",),
    "data_set": ("data_set",),
    "lab_id": ("Lab_ID", "lab_id"),
    "hp_id": ("HP_ID", "hp_id"),
    "test_label": ("Test_label", "test_label"),
    "test_cond": ("test_cond",),
    "flow_config": ("flow_config",),
    "profile_id": ("profile_id",),
    "condition_set_id": ("condition_set_id",),
    "cop_dataset": ("COP_dataset", "cop_dataset", "COP_Dataset"),
    "start_time_s": ("start_time",),
    "end_time_s": ("end_time",),
    "cycle_type": ("cycle_type",),
    "avg_t_db": ("avg_t_db",),
    "avg_t_wb": ("avg_t_wb",),
    "avg_t_supply": ("avg_t_supply",),
    "avg_t_return_emu": ("avg_t_return_emu",),
    "avg_t_return_calc": ("avg_t_return_calc",),
    "DTret": ("DTret",),
    "avg_volume_flow": ("avg_volume_flow",),
    "avg_mass_flow": ("avg_mass_flow",),
    "avg_Q_corr": ("avg_heating_capacity_corr",),
    "avg_P_corr": ("avg_power_input_corr",),
    "COP_corr": ("cop_corr",),
    "avg_Q_uncorr": ("avg_heating_capacity_uncorr",),
    "avg_P_uncorr": ("avg_power_input_uncorr",),
    "COP_uncorr": ("cop_uncorr",),
    "avg_Q_bam_corr": ("avg_heating_capacity_bam_corr",),
    "avg_P_bam_corr": ("avg_power_input_bam_corr",),
    "COP_bam_corr": ("cop_bam_corr",),
    "dev_test_condition": ("dev_test_condition",),
    "dev_db_setpoint": ("dev_db_setpoint",),
    "dev_wb_setpoint": ("dev_wb_setpoint",),
    "dev_tsup_setpoint": ("dev_tsup_setpoint",),
    "dev_total_points": ("dev_total_points",),
    "dev_db_outside_count": ("dev_db_outside_count",),
    "dev_db_percentage": ("dev_db_percentage",),
    "dev_wb_outside_count": ("dev_wb_outside_count",),
    "dev_wb_percentage": ("dev_wb_percentage",),
    "dev_tsup_outside_count": ("dev_tsup_outside_count",),
    "dev_tsup_percentage": ("dev_tsup_percentage",),
    "dev_dtreturn_outside_count": ("dev_dtreturn_outside_count",),
    "dev_dtreturn_percentage": ("dev_dtreturn_percentage",),
    "dev_flow_outside_count": ("dev_flow_outside_count",),
    "dev_flow_percentage": ("dev_flow_percentage",),
    "dev_mean_flow_dev_pct": ("dev_mean_flow_dev_pct",),
    "dev_max_db_pos": ("dev_max_db_pos",),
    "dev_max_db_neg": ("dev_max_db_neg",),
    "dev_max_wb_pos": ("dev_max_wb_pos",),
    "dev_max_wb_neg": ("dev_max_wb_neg",),
    "dev_max_tsup_pos": ("dev_max_tsup_pos",),
    "dev_max_tsup_neg": ("dev_max_tsup_neg",),
    "dev_max_dtreturn_pos": ("dev_max_dtreturn_pos",),
    "dev_max_dtreturn_neg": ("dev_max_dtreturn_neg",),
    "dev_max_flow_pos_pct": ("dev_max_flow_pos_pct",),
    "dev_max_flow_neg_pct": ("dev_max_flow_neg_pct",),
    "notes": ("Notes", "notes"),
    **{c: (c,) for c in OPTIONAL_MONITORING_COLUMNS},
}


def export_prefs_path() -> str:
    return _PREFS_PATH


def resolve_export_columns(selected: Optional[Sequence[str]] = None) -> Tuple[str, ...]:
    """Return selected columns in SUMMARY_COLUMNS order. None/empty → recommended default."""
    if not selected:
        return DEFAULT_EXPORT_COLUMNS
    wanted = {str(c).strip() for c in selected if str(c).strip()}
    resolved = tuple(c for c in SUMMARY_COLUMNS if c in wanted)
    return resolved if resolved else DEFAULT_EXPORT_COLUMNS


def load_default_columns() -> List[str]:
    """Local saved default column list; falls back to recommended default (no empty pass_*)."""
    if not os.path.isfile(_PREFS_PATH):
        return list(DEFAULT_EXPORT_COLUMNS)
    try:
        with open(_PREFS_PATH, "r", encoding="utf-8") as fh:
            data = json.load(fh)
        cols = data.get("default_columns")
        if isinstance(cols, list):
            return list(resolve_export_columns(cols))
    except Exception:
        pass
    return list(DEFAULT_EXPORT_COLUMNS)


def save_default_columns(selected: Sequence[str]) -> List[str]:
    """Persist default column selection (local prefs; not the schema catalogue)."""
    resolved = list(resolve_export_columns(selected))
    os.makedirs(os.path.dirname(_PREFS_PATH), exist_ok=True)
    payload = {
        "schema_id": SCHEMA_ID,
        "default_columns": resolved,
        "updated_at_utc": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
    }
    with open(_PREFS_PATH, "w", encoding="utf-8") as fh:
        json.dump(payload, fh, indent=2, ensure_ascii=False)
        fh.write("\n")
    return resolved


def reset_default_columns() -> List[str]:
    """Remove local prefs so the default becomes the recommended column set again."""
    if os.path.isfile(_PREFS_PATH):
        os.remove(_PREFS_PATH)
    return list(DEFAULT_EXPORT_COLUMNS)


def ui_column_groups() -> Tuple[Tuple[str, Tuple[str, ...]], ...]:
    """Column groups for the Export advanced panel (excludes deferred pass_*)."""
    return COLUMN_GROUPS


def get_tool_version(config: Optional[Dict[str, Any]] = None) -> str:
    """Resolve tool_version for RP-03 (config override → VERSION file → git → fallback)."""
    if config:
        override = (config.get("tool_version") or "").strip()
        if override:
            return override

    version_path = os.path.join(_APP_DIR, "VERSION")
    if os.path.isfile(version_path):
        text = open(version_path, encoding="utf-8").read().strip()
        if text:
            return text

    repo_root = os.path.normpath(os.path.join(_APP_DIR, ".."))
    try:
        out = subprocess.check_output(
            ["git", "rev-parse", "--short", "HEAD"],
            cwd=repo_root,
            stderr=subprocess.DEVNULL,
            text=True,
            timeout=3,
        ).strip()
        if out:
            return f"0.0.0+{out}"
    except Exception:
        pass
    return "0.0.0-dev"


def _norm_key(name: str) -> str:
    return name.lower().replace(" ", "").replace("-", "_")


def _column_map(column_names: Sequence[str]) -> Dict[str, str]:
    """Map export field → actual DB column name present in this database."""
    by_norm = {_norm_key(c): c for c in column_names}
    mapping: Dict[str, str] = {}
    for export_name, candidates in _DB_FIELD_CANDIDATES.items():
        for cand in candidates:
            hit = by_norm.get(_norm_key(cand))
            if hit:
                mapping[export_name] = hit
                break
    return mapping


def _get(row: Any, key: Optional[str]) -> Any:
    if not key:
        return None
    try:
        return row[key]
    except (KeyError, IndexError, TypeError):
        return None


def _appliance_type(profile_id: Optional[str]) -> str:
    if not profile_id:
        return ""
    try:
        import unit_config

        profile = unit_config.get_profile(str(profile_id))
        if profile:
            return str(profile.get("unit_type") or "")
    except Exception:
        pass
    return ""


def db_row_to_summary(
    row: Any,
    *,
    tool_version: str,
    colmap: Dict[str, str],
) -> Dict[str, Any]:
    """Convert one sqlite Row / mapping into a summary dict (None stays None → empty cell)."""
    out: Dict[str, Any] = {name: None for name in SUMMARY_COLUMNS}
    out["tool_version"] = tool_version
    entry_raw = _get(row, "_export_entry_id")
    if entry_raw is None:
        entry_raw = _get(row, "entry_id")
    try:
        out["entry_id"] = int(entry_raw) if entry_raw is not None else None
    except (TypeError, ValueError):
        out["entry_id"] = None

    for export_name, db_col in colmap.items():
        out[export_name] = _get(row, db_col)

    profile_id = out.get("profile_id")
    out["appliance_type"] = _appliance_type(str(profile_id) if profile_id is not None else None)

    # Pass flags reserved for C3
    out["pass_instantaneous"] = None
    out["pass_mean"] = None
    out["pass_controllability"] = None
    return out


def fetch_summary_rows(
    conn: sqlite3.Connection,
    *,
    tool_version: str,
    cop_dataset_only: bool = False,
    cop_values: Sequence[str] = ("YES",),
) -> List[Dict[str, Any]]:
    """Load all results rows (optionally filtered by COP_dataset) as summary dicts."""
    colnames = [r[1] for r in conn.execute("PRAGMA table_info(results)")]
    colmap = _column_map(colnames)
    cop_db = colmap.get("cop_dataset")

    sql = "SELECT rowid AS _export_entry_id, * FROM results"
    params: List[Any] = []
    if cop_dataset_only:
        if not cop_db:
            return []
        wanted = [str(v).strip().upper() for v in cop_values if str(v).strip()]
        if not wanted:
            wanted = ["YES"]
        placeholders = ",".join("?" for _ in wanted)
        sql += (
            f' WHERE upper(trim(cast("{cop_db}" AS TEXT))) IN ({placeholders})'
        )
        params.extend(wanted)
    sql += " ORDER BY _export_entry_id"

    rows = conn.execute(sql, params).fetchall()
    return [db_row_to_summary(r, tool_version=tool_version, colmap=colmap) for r in rows]


def _cell_for_csv(value: Any) -> str:
    if value is None:
        return ""
    return value


def rows_to_csv_text(
    rows: Iterable[Dict[str, Any]],
    columns: Optional[Sequence[str]] = None,
) -> str:
    cols = resolve_export_columns(columns)
    buf = StringIO()
    writer = csv.DictWriter(buf, fieldnames=list(cols), lineterminator="\n", extrasaction="ignore")
    writer.writeheader()
    for row in rows:
        writer.writerow({k: _cell_for_csv(row.get(k)) for k in cols})
    return buf.getvalue()


def rows_to_xlsx_bytes(
    rows: Sequence[Dict[str, Any]],
    *,
    tool_version: str,
    cop_dataset_only: bool,
    columns: Optional[Sequence[str]] = None,
) -> bytes:
    from openpyxl import Workbook

    cols = resolve_export_columns(columns)
    wb = Workbook()
    ws = wb.active
    ws.title = "summary"
    ws.append(list(cols))
    for row in rows:
        ws.append([_cell_for_csv(row.get(k)) for k in cols])

    meta = wb.create_sheet("meta")
    meta.append(["key", "value"])
    meta.append(["schema_id", SCHEMA_ID])
    meta.append(["tool_version", tool_version])
    meta.append(["cop_dataset_only", "yes" if cop_dataset_only else "no"])
    meta.append(["column_count", len(cols)])
    meta.append(["columns_custom", "yes" if cols != DEFAULT_EXPORT_COLUMNS else "no"])
    meta.append(["row_count", len(rows)])
    meta.append(["exported_at_utc", datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")])
    meta.append(["columns", ",".join(cols)])

    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio.read()


def build_export_filename(
    *,
    fmt: str,
    cop_dataset_only: bool,
    tool_version: str,
    columns: Optional[Sequence[str]] = None,
) -> str:
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    scope = "copYES" if cop_dataset_only else "all"
    cols = resolve_export_columns(columns)
    if cols != DEFAULT_EXPORT_COLUMNS:
        scope = f"{scope}_cols{len(cols)}"
    safe_ver = "".join(c if c.isalnum() or c in "._-" else "_" for c in tool_version)
    ext = "xlsx" if fmt == "xlsx" else "csv"
    return f"t42_summary_{scope}_{safe_ver}_{stamp}.{ext}"


def build_export_bytes(
    conn: sqlite3.Connection,
    *,
    fmt: str = "csv",
    tool_version: str,
    cop_dataset_only: bool = False,
    cop_values: Sequence[str] = ("YES",),
    columns: Optional[Sequence[str]] = None,
) -> Tuple[bytes, str, str, int]:
    """Return (payload, filename, mimetype, row_count)."""
    cols = resolve_export_columns(columns)
    rows = fetch_summary_rows(
        conn,
        tool_version=tool_version,
        cop_dataset_only=cop_dataset_only,
        cop_values=cop_values,
    )
    filename = build_export_filename(
        fmt=fmt,
        cop_dataset_only=cop_dataset_only,
        tool_version=tool_version,
        columns=cols,
    )
    if fmt == "xlsx":
        payload = rows_to_xlsx_bytes(
            rows,
            tool_version=tool_version,
            cop_dataset_only=cop_dataset_only,
            columns=cols,
        )
        mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    else:
        payload = rows_to_csv_text(rows, columns=cols).encode("utf-8-sig")
        mime = "text/csv; charset=utf-8"
    return payload, filename, mime, len(rows)
